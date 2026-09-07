"""
Carga robusta de .xlsx para dossiers y resultados Grill.

Algunos exports (Grill / Excel no estándar) traen XML de relaciones
(xl/**/_rels/*.rels) mal formado o enorme. openpyxl parsea esos .rels
con defusedxml y revienta en get_dependents → fromstring (ParseError).

Estrategia:
  1. openpyxl (conserva hipervínculos de celda).
  2. Si falla: sanitizar XML del paquete OOXML y reintentar openpyxl.
  3. Si los .rels siguen rotos: vaciarlos y reintentar.
  4. Último recurso: python-calamine / pandas y reconstruir un Workbook
     mínimo para el resto del pipeline.
"""
from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path
from xml.etree.ElementTree import ParseError as ETParseError
from xml.etree.ElementTree import fromstring as xml_fromstring

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

EMPTY_RELS = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'
)

# Entidades XML permitidas; el resto de "&" sueltos se escapan.
_BARE_AMP = re.compile(r"&(?!(?:amp|lt|gt|apos|quot|#\d+|#x[0-9A-Fa-f]+);)")
_ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

PREFERRED_SHEETS = ("Table1", "Resultado", "Sheet1")

XLSX_READ_ERROR = (
    "No se pudo leer el archivo Excel. Suele ocurrir con resultados Grill "
    "u otros .xlsx con XML de relaciones no estándar. Pruebe re-guardar el "
    "archivo en Excel como .xlsx, o use Python 3.11/3.12 en Streamlit Cloud."
)


def _read_bytes(source) -> bytes:
    if isinstance(source, (bytes, bytearray, memoryview)):
        return bytes(source)
    if hasattr(source, "read"):
        if hasattr(source, "seek"):
            try:
                source.seek(0)
            except Exception:
                pass
        data = source.read()
        if hasattr(source, "seek"):
            try:
                source.seek(0)
            except Exception:
                pass
        if isinstance(data, str):
            data = data.encode("utf-8")
        return data
    return Path(source).read_bytes()


def _should_fallback(exc: BaseException) -> bool:
    if isinstance(exc, (OSError, InvalidFileException, ETParseError, zipfile.BadZipFile)):
        return True
    name = type(exc).__name__
    if name in {"ParseError", "XMLSyntaxError"}:
        return True
    # styles.xml no estándar
    if isinstance(exc, (IndexError, KeyError)):
        return True
    return False


def _sanitize_xml_bytes(raw: bytes) -> bytes:
    text = raw.decode("utf-8", errors="replace").lstrip("\ufeff")
    text = _ILLEGAL_XML_CHARS.sub("", text)
    text = _BARE_AMP.sub("&amp;", text)
    return text.encode("utf-8")


def _is_xml_part(name: str) -> bool:
    lower = name.lower()
    return lower.endswith(".xml") or lower.endswith(".rels")


def sanitize_xlsx_bytes(data: bytes, *, blank_unparseable_rels: bool = False) -> bytes:
    """Reescribe el zip OOXML sanitizando XML y, opcionalmente, vaciando .rels rotos."""
    in_buf = io.BytesIO(data)
    out_buf = io.BytesIO()
    with zipfile.ZipFile(in_buf, "r") as zin, zipfile.ZipFile(
        out_buf, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for info in zin.infolist():
            content = zin.read(info.filename)
            if _is_xml_part(info.filename):
                try:
                    xml_fromstring(content)
                except Exception:
                    content = _sanitize_xml_bytes(content)
                    if blank_unparseable_rels and info.filename.lower().endswith(".rels"):
                        try:
                            xml_fromstring(content)
                        except Exception:
                            content = EMPTY_RELS
            zout.writestr(info.filename, content)
    return out_buf.getvalue()


def _cell_value(val):
    if val is None:
        return None
    try:
        if pd.isna(val):
            return None
    except (TypeError, ValueError):
        pass
    return val


def workbook_from_tabular_bytes(data: bytes) -> Workbook:
    """Lee celdas con calamine (o pandas) y arma un Workbook openpyxl mínimo."""
    buf = io.BytesIO(data)
    last_err = None
    raw = None
    try:
        raw = pd.read_excel(buf, header=None, engine="calamine")
    except Exception as exc:
        last_err = exc
        buf.seek(0)
        try:
            raw = pd.read_excel(buf, header=None, engine="openpyxl")
        except Exception as exc2:
            last_err = exc2

    if raw is None or raw.empty:
        detail = f" ({last_err})" if last_err else ""
        raise RuntimeError(
            "No se pudo leer el contenido del Excel con el motor alternativo "
            f"(calamine/pandas).{detail}"
        )

    wb = Workbook()
    ws = wb.active
    for r_idx, row in enumerate(raw.itertuples(index=False), start=1):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=_cell_value(val))
    return wb


def load_xlsx_workbook(source, *, data_only: bool = False) -> Workbook:
    """
    Carga un .xlsx.

    Devuelve un Workbook openpyxl. Si el archivo tiene XML de relaciones
    roto, intenta sanitizarlo y, si hace falta, reconstruye el libro desde
    calamine. Si nada funciona, lanza RuntimeError con mensaje en español.
    """
    data = _read_bytes(source)
    if not data:
        raise RuntimeError("El archivo Excel está vacío.")

    errors = []

    try:
        return load_workbook(io.BytesIO(data), data_only=data_only)
    except Exception as exc:
        if not _should_fallback(exc):
            raise
        errors.append(exc)

    try:
        sanitized = sanitize_xlsx_bytes(data, blank_unparseable_rels=False)
        return load_workbook(io.BytesIO(sanitized), data_only=data_only)
    except Exception as exc:
        errors.append(exc)

    try:
        sanitized = sanitize_xlsx_bytes(data, blank_unparseable_rels=True)
        return load_workbook(io.BytesIO(sanitized), data_only=data_only)
    except Exception as exc:
        errors.append(exc)

    try:
        return workbook_from_tabular_bytes(data)
    except Exception as exc:
        errors.append(exc)

    last = errors[-1] if errors else None
    raise RuntimeError(f"{XLSX_READ_ERROR} Detalle: {last}") from last


def read_excel_sheets(source) -> dict:
    """Lee todas las hojas como DataFrames (configuración u otros .xlsx)."""
    data = _read_bytes(source)
    if not data:
        raise RuntimeError("El archivo Excel está vacío.")

    errors = []

    try:
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, engine="openpyxl")
        if sheets:
            return sheets
    except Exception as exc:
        if not _should_fallback(exc):
            raise
        errors.append(exc)

    try:
        sanitized = sanitize_xlsx_bytes(data, blank_unparseable_rels=True)
        sheets = pd.read_excel(io.BytesIO(sanitized), sheet_name=None, engine="openpyxl")
        if sheets:
            return sheets
    except Exception as exc:
        errors.append(exc)

    try:
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, engine="calamine")
        if sheets:
            return sheets
    except Exception as exc:
        errors.append(exc)

    last = errors[-1] if errors else None
    raise RuntimeError(f"{XLSX_READ_ERROR} Detalle: {last}") from last


def pick_worksheet(wb: Workbook):
    """Prefiere Table1 (exports Grill) y, si no, la hoja activa."""
    names_lower = {n.lower(): n for n in wb.sheetnames}
    for cand in PREFERRED_SHEETS:
        real = names_lower.get(cand.lower())
        if real:
            return wb[real]
    return wb.active


def _is_blank(val) -> bool:
    if val is None:
        return True
    try:
        if pd.isna(val):
            return True
    except (TypeError, ValueError):
        pass
    if isinstance(val, str) and not val.strip():
        return True
    return False


def iter_sheet_records(sheet):
    """
    Recorre filas de datos alineadas por índice de columna.

    No descarta una fila solo porque las primeras columnas (p. ej.
    «Tipo de nota» / «Empresas Consulta») vengan vacías. Solo se omiten
    filas sin ningún valor en columnas con encabezado.
    """
    header_row = [c.value for c in sheet[1]]
    named_idx = [
        (i, h if not isinstance(h, str) else h.strip())
        for i, h in enumerate(header_row)
        if h is not None and str(h).strip() != ""
    ]
    if not named_idx:
        return

    for row in sheet.iter_rows(min_row=2):
        values = [c.value for c in row]
        picked = []
        cells_by_header = {}
        record = {}
        for i, header in named_idx:
            val = values[i] if i < len(values) else None
            record[header] = val
            cells_by_header[header] = row[i] if i < len(row) else None
            picked.append(val)
        if all(_is_blank(v) for v in picked):
            continue
        yield record, cells_by_header
