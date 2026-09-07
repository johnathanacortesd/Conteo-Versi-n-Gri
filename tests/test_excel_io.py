"""Carga de xlsx: happy path openpyxl + fallback ante .rels rotos."""
from __future__ import annotations

import io
import unittest
import zipfile
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch
from xml.etree.ElementTree import ParseError

from openpyxl import Workbook, load_workbook

from dossier import load_config, process_dossier
from excel_io import (
    iter_sheet_records,
    load_xlsx_workbook,
    pick_worksheet,
    read_excel_sheets,
    workbook_from_tabular_bytes,
)

ROOT = Path(__file__).resolve().parents[1]


def _dossier_xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "NoticiaId",
            "Fecha",
            "Medio",
            "Tipo de Medio",
            "Título",
            "Link Nota",
            "Empresa rel.",
        ]
    )
    ws.append(
        [
            101,
            date(2026, 1, 15),
            "El Tiempo",
            "online",
            "Nissan lanza modelo",
            "ver nota",
            "Nissan",
        ]
    )
    ws.cell(row=2, column=6).hyperlink = "https://example.com/nota"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _corrupt_worksheet_rels(xlsx_bytes: bytes, payload: bytes | None = None) -> bytes:
    """Inyecta XML inválido en xl/worksheets/_rels/*.rels (el crash de get_dependents)."""
    in_buf = io.BytesIO(xlsx_bytes)
    out_buf = io.BytesIO()
    with zipfile.ZipFile(in_buf, "r") as zin, zipfile.ZipFile(out_buf, "w") as zout:
        found = False
        for info in zin.infolist():
            content = zin.read(info.filename)
            if "worksheets/_rels/" in info.filename.replace("\\", "/") and info.filename.endswith(
                ".rels"
            ):
                found = True
                if payload is not None:
                    content = payload
                else:
                    text = content.decode("utf-8")
                    broken = (
                        '<Relationship Id="rIdBad" '
                        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
                        'Target="https://example.com/a?b=1&c=2" TargetMode="External"/>'
                    )
                    if "</Relationships>" in text:
                        text = text.replace("</Relationships>", broken + "</Relationships>")
                    else:
                        text = broken
                    content = text.encode("utf-8")
            zout.writestr(info.filename, content)
        if not found:
            zout.writestr(
                "xl/worksheets/_rels/sheet1.xml.rels",
                payload
                or (
                    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    b'<Relationship Id="rIdBad" Type="http://example.com/rel" '
                    b'Target="https://example.com/a?b=1&c=2"/></Relationships>'
                ),
            )
    return out_buf.getvalue()


class LoadXlsxWorkbookTests(unittest.TestCase):
    def test_happy_path_openpyxl_file_still_loads(self):
        data = _dossier_xlsx_bytes()
        wb = load_xlsx_workbook(io.BytesIO(data))
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "NoticiaId")
        self.assertEqual(rows[1][0], 101)
        self.assertEqual(rows[1][4], "Nissan lanza modelo")
        link_cell = wb.active.cell(row=2, column=6)
        self.assertEqual(link_cell.value, "ver nota")
        self.assertTrue(
            link_cell.hyperlink and "example.com/nota" in str(link_cell.hyperlink.target)
        )

    def test_broken_rels_crashes_openpyxl_but_fallback_loads(self):
        broken = _corrupt_worksheet_rels(_dossier_xlsx_bytes())

        with self.assertRaises(ParseError):
            load_workbook(io.BytesIO(broken))

        wb = load_xlsx_workbook(io.BytesIO(broken))
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][5], "Link Nota")
        self.assertEqual(rows[1][4], "Nissan lanza modelo")
        self.assertEqual(rows[1][6], "Nissan")
        self.assertEqual(wb.active.cell(row=2, column=6).value, "ver nota")

    def test_garbage_rels_still_loads_via_sanitize_or_calamine(self):
        broken = _corrupt_worksheet_rels(
            _dossier_xlsx_bytes(),
            payload=b"this is not xml &&& <Relationships",
        )

        with self.assertRaises(Exception):
            load_workbook(io.BytesIO(broken))

        wb = load_xlsx_workbook(io.BytesIO(broken))
        titles = [c.value for c in wb.active[1]]
        self.assertIn("Título", titles)

    def test_calamine_rebuild_when_openpyxl_always_fails(self):
        data = _dossier_xlsx_bytes()

        def boom(*_args, **_kwargs):
            raise ParseError("not well-formed (invalid token)")

        with patch("excel_io.load_workbook", side_effect=boom):
            wb = load_xlsx_workbook(io.BytesIO(data))

        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "NoticiaId")
        self.assertEqual(rows[1][4], "Nissan lanza modelo")

    def test_spanish_error_when_bytes_are_not_xlsx(self):
        with self.assertRaises(RuntimeError) as ctx:
            load_xlsx_workbook(io.BytesIO(b"esto no es un excel"))
        self.assertIn("No se pudo leer el archivo Excel", str(ctx.exception))

    def test_workbook_from_tabular_bytes_roundtrip(self):
        wb = workbook_from_tabular_bytes(_dossier_xlsx_bytes())
        values = [c.value for c in next(wb.active.iter_rows(min_row=2, max_row=2))]
        self.assertEqual(values[3], "online")
        self.assertTrue(
            values[1] == date(2026, 1, 15)
            or isinstance(values[1], datetime)
            or str(values[1]).startswith("2026-01-15")
        )

    def test_read_excel_sheets_config_happy_path(self):
        cfg = ROOT / "Configuracion.xlsx"
        if not cfg.exists():
            self.skipTest("Configuracion.xlsx no está en el checkout")
        sheets = read_excel_sheets(cfg)
        self.assertIn("Regiones", sheets)
        self.assertIn("Internet", sheets)
        self.assertGreater(len(sheets["Regiones"]), 10)

    def test_read_excel_sheets_broken_rels(self):
        broken = _corrupt_worksheet_rels(_dossier_xlsx_bytes())
        sheets = read_excel_sheets(io.BytesIO(broken))
        self.assertTrue(sheets)
        first = next(iter(sheets.values()))
        self.assertIn("NoticiaId", list(first.columns) + list(first.iloc[0].values))


class SheetRecordTests(unittest.TestCase):
    def test_pick_table1_over_active(self):
        wb = Workbook()
        wb.active.title = "Cover"
        ws = wb.create_sheet("Table1")
        ws.append(["Tipo de Medio", "Título"])
        ws.append(["Internet", "Nota"])
        self.assertEqual(pick_worksheet(wb).title, "Table1")

    def test_keeps_rows_when_early_grill_columns_are_blank(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Table1"
        ws.append(
            [
                "Tipo de nota",
                "Empresas Consulta",
                "Tipo de Medio",
                "Título",
                "Link (Streaming - Imagen)",
                "Tono",
                "Subtema",
            ]
        )
        # Primera fila de datos: columnas tempranas vacías, el resto con nota.
        ws.append([None, "", "Internet", "Chery en feria", "https://example.com/a", "Positivo", "Lanzamiento"])
        ws.append(["Informativa", "Chery", "Radio", "Entrevista", None, "Neutro", "Producto"])
        records = list(iter_sheet_records(ws))
        self.assertEqual(len(records), 2)
        self.assertEqual(records[0][0]["Título"], "Chery en feria")
        self.assertEqual(records[0][0]["Tipo de Medio"], "Internet")
        self.assertTrue(
            records[0][0]["Tipo de nota"] in (None, "")
            or str(records[0][0]["Tipo de nota"]).strip() == ""
        )
        self.assertEqual(records[1][0]["Tipo de Medio"], "Radio")

    def test_blank_header_cells_do_not_shift_columns(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value=None)
        ws.cell(row=1, column=2, value="Tipo de Medio")
        ws.cell(row=1, column=3, value="Título")
        ws.cell(row=2, column=1, value="basura")
        ws.cell(row=2, column=2, value="Prensa")
        ws.cell(row=2, column=3, value="Titular alineado")
        records = list(iter_sheet_records(ws))
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0][0]["Tipo de Medio"], "Prensa")
        self.assertEqual(records[0][0]["Título"], "Titular alineado")
        self.assertNotIn(None, records[0][0])


class ProcessDossierTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cfg = ROOT / "Configuracion.xlsx"
        cls.rmap, cls.imap = load_config(cfg) if cfg.exists() else ({}, {})

    def _maps(self):
        return self.rmap or {"el tiempo": "Nacional"}, self.imap or {}

    def test_raw_dossier_counts_av_and_grafica(self):
        data = _dossier_xlsx_bytes()
        df, av, gr = process_dossier(io.BytesIO(data), *self._maps())
        self.assertEqual(av, 0)
        self.assertEqual(gr, 1)
        self.assertEqual(df.iloc[0]["Tipo de Medio"], "Internet")
        self.assertEqual(df.iloc[0]["Título"], "Nissan lanza modelo")
        self.assertEqual(df.iloc[0]["Link Nota"], "https://example.com/nota")

    def test_broken_rels_dossier_still_processed(self):
        broken = _corrupt_worksheet_rels(_dossier_xlsx_bytes())
        df, av, gr = process_dossier(io.BytesIO(broken), *self._maps())
        self.assertEqual(gr, 1)
        self.assertEqual(df.iloc[0]["Título"], "Nissan lanza modelo")

    def test_grill_export_with_blank_early_columns(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Table1"
        ws.append(
            [
                "Tipo de nota",
                "Empresas Consulta",
                "ID Noticia",
                "Fecha",
                "Medio",
                "Tipo de Medio",
                "Título",
                "Tono",
                "Tema",
                "Subtema",
                "Link Nota",
                "Link (Streaming - Imagen)",
                "Menciones - Empresa",
                "Resumen - Aclaracion",
            ]
        )
        ws.append(
            [
                None,
                None,
                55,
                date(2026, 3, 1),
                "El Tiempo",
                "Internet",
                "Ya clasificada",
                "Positivo",
                "Producto",
                "Lanzamiento",
                "https://example.com/g1",
                "https://example.com/img",
                "Chery",
                "Resumen grill",
            ]
        )
        ws.append(
            [
                "",
                "",
                56,
                date(2026, 3, 2),
                "Caracol",
                "Radio",
                "Nota AV grill",
                "Neutro",
                "Marca",
                None,
                "https://example.com/av",
                None,
                "Nissan",
                "Cuerpo AV",
            ]
        )
        buf = io.BytesIO()
        wb.save(buf)
        df, av, gr = process_dossier(io.BytesIO(buf.getvalue()), *self._maps())
        self.assertEqual(len(df), 2)
        self.assertEqual(av, 1)
        self.assertEqual(gr, 1)
        self.assertEqual(df.iloc[0]["Tema"], "Producto")
        self.assertEqual(df.iloc[0]["Resumen - Aclaracion"], "Resumen grill")
        self.assertEqual(df.iloc[0]["Link (Streaming - Imagen)"], "https://example.com/img")
        self.assertEqual(df.iloc[1]["Link Nota"], "https://example.com/av")


if __name__ == "__main__":
    unittest.main()
